#!/usr/bin/env python3
"""Build hunting/checklists/IR_CHECKLIST_ROLES.xlsx from the Markdown role checklists.

One sheet per checklist, laid out like the Overall Status sheet of the IR checklist
(IR_CHECKLIST_ED.xlsx): item number and check in columns A and B, then the tracking
columns Tasked to, Status, Findings, Handover From, Handover to, Handover Date and Details.
Details holds "What to look for"; the query IDs and the IR checklist reference follow
in two extra columns.

Each checklist must keep the Markdown layout described in hunting/checklists/README.md:
a "# Checklist: <name>" title, "## <n>. <section>" headings each followed by one table
with the columns #, Check, Queries, IR ref, What to look for, and an optional
"## Not covered by these queries" bullet list.

Usage:
    pip install openpyxl
    python3 hunting/tools/build_checklist_xlsx.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

HERE = Path(__file__).resolve().parent
CHECKLISTS = HERE.parent / "checklists"
OUTPUT = CHECKLISTS / "IR_CHECKLIST_ROLES.xlsx"

ORDER = [
    "00-baseline-all-windows-hosts.md",
    "end-user-workstation.md",
    "web-server.md",
    "database-server.md",
    "domain-controller.md",
    "file-server.md",
]
SHEET_NAMES = {
    "00-baseline-all-windows-hosts.md": "Baseline",
    "end-user-workstation.md": "Workstation",
    "web-server.md": "Web server",
    "database-server.md": "Database server",
    "domain-controller.md": "Domain controller",
    "file-server.md": "File server",
}
HEADERS = ["#", "Check", "Tasked to", "Status", "Findings", "Handover From",
           "Handover to", "Handover Date", "Details", "Queries", "IR ref"]
WIDTHS = [7, 55, 12, 13, 30, 14, 12, 14, 60, 30, 14]
STATUSES = "Not started,In progress,Done,Not applicable,No data"

NAVY = "FF002060"
SECTION_FILL = PatternFill("solid", fgColor="FFD9E1F2")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
THIN = Side(style="thin", color="FF808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]*\)")


def plain(md: str) -> str:
    """Markdown cell text to plain text: links become their label, code ticks and bold go."""
    text = LINK_RE.sub(r"\1", md)
    return text.replace("`", "").replace("**", "").strip()


def split_row(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def parse(path: Path):
    """Return (title, intro, sections, not_covered); sections = [(heading, [row cells])]."""
    lines = path.read_text(encoding="utf-8").splitlines()
    title = intro = ""
    sections: list[tuple[str, list[list[str]]]] = []
    not_covered: list[str] = []
    current = None
    in_not_covered = False
    for line in lines:
        if line.startswith("# "):
            title = line[2:].replace("Checklist:", "").strip()
        elif line.startswith("## "):
            heading = line[3:].strip()
            in_not_covered = heading.lower().startswith("not covered")
            current = None if in_not_covered else (heading, [])
            if current:
                sections.append(current)
        elif in_not_covered and line.startswith("- "):
            not_covered.append(plain(line[2:]))
        elif current is not None and line.startswith("|"):
            cells = split_row(line)
            if cells[0] in ("#", "") or set(cells[0]) <= set("-:"):
                continue
            if len(cells) != 5:
                sys.exit(f"{path.name}: table row does not have 5 columns: {line}")
            current[1].append(cells)
        elif not title or sections or not line.strip():
            continue
        elif not intro:
            intro = plain(line)
    if not sections:
        sys.exit(f"{path.name}: no sections found")
    return title, intro, sections, not_covered


def style_row(ws, row: int, *, bold=False, fill=None, font_color=None):
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=row, column=col)
        cell.border = BORDER
        cell.alignment = WRAP
        cell.font = Font(bold=bold, color=font_color)
        if fill:
            cell.fill = fill


def build_sheet(wb, path: Path):
    title, intro, sections, not_covered = parse(path)
    ws = wb.create_sheet(SHEET_NAMES.get(path.name, path.stem[:31]))
    for i, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

    ws.cell(row=1, column=1, value=f"Project <name>  {title} checklist").font = Font(bold=True, size=20, color=NAVY)
    ws.row_dimensions[1].height = 25.5
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws.cell(row=2, column=1, value=intro).alignment = WRAP
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS))
    ws.row_dimensions[2].height = 60

    for col, name in enumerate(HEADERS, start=1):
        ws.cell(row=3, column=col, value=name)
    style_row(ws, 3, bold=True, fill=HEADER_FILL, font_color="FFFFFFFF")
    ws.freeze_panes = "C4"

    status = DataValidation(type="list", formula1=f'"{STATUSES}"', allow_blank=True)
    ws.add_data_validation(status)

    row = 4
    for heading, items in sections:
        num, _, name = heading.partition(" ")
        ws.cell(row=row, column=1, value=num.rstrip("."))
        ws.cell(row=row, column=2, value=name)
        style_row(ws, row, bold=True, fill=SECTION_FILL)
        row += 1
        for item, check, queries, ir_ref, look_for in items:
            values = {1: item, 2: plain(check), 9: plain(look_for), 10: plain(queries), 11: plain(ir_ref)}
            for col, value in values.items():
                ws.cell(row=row, column=col, value=value)
            style_row(ws, row)
            status.add(ws.cell(row=row, column=4))
            row += 1
    if not_covered:
        row += 1
        ws.cell(row=row, column=2, value="Not covered by these queries")
        style_row(ws, row, bold=True, fill=SECTION_FILL)
        row += 1
        for text in not_covered:
            ws.cell(row=row, column=1, value="*")
            ws.cell(row=row, column=2, value=text)
            style_row(ws, row)
            status.add(ws.cell(row=row, column=4))
            row += 1
    ws.auto_filter.ref = f"A3:{ws.cell(row=3, column=len(HEADERS)).column_letter}{row - 1}"


def main():
    files = [CHECKLISTS / name for name in ORDER if (CHECKLISTS / name).exists()]
    extra = sorted(p for p in CHECKLISTS.glob("*.md") if p.name not in ORDER and p.name != "README.md")
    wb = Workbook()
    wb.remove(wb.active)
    for path in files + extra:
        build_sheet(wb, path)
    wb.save(OUTPUT)
    print(f"wrote {OUTPUT.relative_to(HERE.parent.parent)} ({len(wb.sheetnames)} sheets)")


if __name__ == "__main__":
    main()
